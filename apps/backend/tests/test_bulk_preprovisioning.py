"""BP-01 fase B: carga masiva de la estructura (organizaciones + miembros + módulos),
vista previa, envío de invitaciones, reenvío de vencidas y descarte."""
from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import delete, func, select

from hg.db import SessionLocal
from hg.modules.company.bulk_service import (
    MODULE_HEADERS,
    ORG_HEADERS,
    TEMPLATE_HEADERS,
    XLSX_MIME,
)
from hg.modules.identity.invitations import Invitation
from hg.modules.identity.models import Organization, User, UserRole
from hg.modules.learning_units.models import LearningUnit, ModuleAssignment, OrgModuleAssignment

API = "/api/v1"
SEND = "hg.modules.identity.service.email_service.send"


def _email() -> str:
    return f"p-{uuid4().hex[:10]}@hgtest.test"


def _book(orgs=None, members=None, modules=None) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if orgs is not None:
        ws = wb.create_sheet("Organizaciones")
        ws.append(ORG_HEADERS)
        for r in orgs:
            ws.append(r)
    if members is not None:
        ws = wb.create_sheet("Miembros")
        ws.append(TEMPLATE_HEADERS)
        for r in members:
            ws.append(r)
    if modules is not None:
        ws = wb.create_sheet("Módulos")
        ws.append(MODULE_HEADERS)
        for r in modules:
            ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, headers, content, **params):
    return client.post(
        f"{API}/company/members/bulk-import",
        headers=headers,
        params=params or None,
        files={"file": ("estructura.xlsx", content, XLSX_MIME)},
    )


def _setup(factory, licenses_total: int = 50):
    org = factory.make_org(slug=f"co-{uuid4().hex[:8]}", licenses_total=licenses_total)
    admin = factory.make_user(org=org, role=UserRole.company_admin)
    return org, admin


def _members(client, headers) -> dict[str, dict]:
    return {m["email"]: m for m in client.get(f"{API}/company/members", headers=headers).json()}


def _make_unit(slug: str) -> uuid4:
    s = SessionLocal()
    try:
        unit = LearningUnit(
            slug=slug, title="Módulo de prueba", dimension_code="CP", level_code="L1",
            pillar_code="P1", published_at=datetime.now(UTC),
        )
        s.add(unit)
        s.commit()
        return unit.id
    finally:
        s.close()


def _drop_unit(unit_id) -> None:
    s = SessionLocal()
    try:
        s.execute(delete(OrgModuleAssignment).where(OrgModuleAssignment.learning_unit_id == unit_id))
        s.execute(delete(ModuleAssignment).where(ModuleAssignment.learning_unit_id == unit_id))
        s.execute(delete(LearningUnit).where(LearningUnit.id == unit_id))
        s.commit()
    finally:
        s.close()


def test_three_sheets_create_orgs_invited_members_and_link_managers(
    client: TestClient, factory, auth_headers
) -> None:
    _org, admin = _setup(factory, licenses_total=50)
    boss, rep = _email(), _email()
    content = _book(
        orgs=[["Ventas Norte", "", 10]],
        # el reporte aparece ANTES que su manager: se enlaza en la 2ª pasada
        members=[
            ["ventas-norte", rep, "Rita Reporte", "collaborator", boss, "Analista", "Comercial"],
            ["Ventas Norte", boss, "Beto Jefe", "manager", "", "", ""],
        ],
    )
    with patch(SEND) as send:
        res = _post(client, auth_headers(admin), content)
    assert res.status_code == 200, res.text
    body = res.json()
    assert [o["estado"] for o in body["organizaciones"]] == ["creada"]
    assert body["creados"] == 2 and body["errores"] == 0
    assert all(r.get("aviso") is None for r in body["filas"])
    send.assert_not_called()  # cargar la estructura NO manda correos

    members = _members(client, auth_headers(admin))
    assert members[rep]["is_invited"] and members[rep]["invite_status"] == "staged"
    assert members[rep]["manager_name"] == "Beto Jefe"  # manager también invitado
    assert members[boss]["role"] == "manager"

    orgs = {o["name"]: o for o in client.get(f"{API}/company/organizations", headers=auth_headers(admin)).json()}
    assert orgs["Ventas Norte"]["license_quota"] == 10

    factory.session.expire_all()
    user = factory.session.scalar(select(User).where(func.lower(User.email) == rep))
    assert user.job_title == "Analista" and user.department == "Comercial"


def test_validate_only_writes_nothing_and_sends_nothing(
    client: TestClient, factory, auth_headers
) -> None:
    org, admin = _setup(factory)
    email = _email()
    content = _book(
        orgs=[["Preview Org", "", 5]],
        members=[["Preview Org", email, "Pre Vista", "collaborator", "", "", ""]],
    )
    with patch(SEND) as send:
        res = _post(client, auth_headers(admin), content, validate_only=True, send_invitations=True)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["solo_validar"] is True
    assert body["organizaciones"][0]["estado"] == "creada"
    assert body["creados"] == 1 and body["errores"] == 0
    send.assert_not_called()

    factory.session.expire_all()
    assert factory.session.scalar(select(User).where(func.lower(User.email) == email)) is None
    assert (
        factory.session.scalar(
            select(func.count()).select_from(Organization).where(
                Organization.company_id == org.company_id, Organization.name == "Preview Org"
            )
        )
        == 0
    )


def test_manager_must_exist_in_file_or_company(client: TestClient, factory, auth_headers) -> None:
    _org, admin = _setup(factory)
    collab = _email()
    content = _book(
        orgs=[["Equipo X", "", 0]],
        members=[
            ["Equipo X", _email(), "A", "collaborator", _email(), "", ""],  # no existe
            ["Equipo X", collab, "B", "collaborator", "", "", ""],
            ["Equipo X", _email(), "C", "collaborator", collab, "", ""],  # no es manager
        ],
    )
    body = _post(client, auth_headers(admin), content).json()
    assert body["errores"] == 2
    motivos = " ".join(r["motivo"] for r in body["filas"] if r["motivo"])
    assert "no existe" in motivos and "no es manager" in motivos


def test_org_quota_cannot_exceed_company_pool(client: TestClient, factory, auth_headers) -> None:
    _org, admin = _setup(factory, licenses_total=10)
    content = _book(orgs=[["Cupo Grande", "", 500]])
    body = _post(client, auth_headers(admin), content).json()
    assert body["organizaciones"][0]["estado"] == "error"
    assert "pool" in body["organizaciones"][0]["motivo"].lower() or "cupo" in body["organizaciones"][0]["motivo"].lower()


def test_modules_sheet_assigns_to_org_and_invited_receive_them_on_accept(
    client: TestClient, factory, auth_headers
) -> None:
    _org, admin = _setup(factory)
    slug = f"bp-{uuid4().hex[:8]}"
    unit_id = _make_unit(slug)
    try:
        email = _email()
        content = _book(
            orgs=[["Con Modulos", "", 0]],
            members=[["Con Modulos", email, "Nuevo Ingreso", "collaborator", "", "", ""]],
            modules=[["Con Modulos", slug, "2099-12-31"], ["Con Modulos", "no-existe-xyz", ""]],
        )
        body = _post(client, auth_headers(admin), content).json()
        assert [m["estado"] for m in body["modulos"]] == ["asignado", "error"]
        assert "no existe" in body["modulos"][1]["motivo"]

        # Envía y acepta: ya entra con el módulo asignado.
        members = _members(client, auth_headers(admin))
        uid = members[email]["id"]
        with patch(SEND):
            sent = client.post(
                f"{API}/company/invitations/send", headers=auth_headers(admin), json={}
            )
        assert sent.status_code == 200 and sent.json()["enviadas"] >= 1
        factory.session.expire_all()
        invitation = factory.session.scalar(
            select(Invitation).where(func.lower(Invitation.email) == email)
        )
        assert invitation is not None
        # El token plano solo existe en el correo: se emite uno conocido para aceptar.
        from hg.core.security import generate_opaque_token

        plain, token_hash = generate_opaque_token()
        invitation.token_hash = token_hash
        factory.session.commit()
        acc = client.post(
            "/api/v1/auth/accept-invite",
            json={"token": plain, "password": "Brand0New!!", "username_or_email": "nuevo" + uuid4().hex[:4]},
        )
        assert acc.status_code == 200, acc.text
        factory.session.expire_all()
        assigned = factory.session.scalars(
            select(ModuleAssignment.learning_unit_id).where(ModuleAssignment.user_id == uid)
        ).all()
        assert unit_id in assigned
    finally:
        _drop_unit(unit_id)


def test_send_is_explicit_and_not_repeated(client: TestClient, factory, auth_headers) -> None:
    _org, admin = _setup(factory)
    a, b = _email(), _email()
    content = _book(
        orgs=[["Envios", "", 0]],
        members=[
            ["Envios", a, "A", "collaborator", "", "", ""],
            ["Envios", b, "B", "collaborator", "", "", ""],
        ],
    )
    _post(client, auth_headers(admin), content)
    with patch(SEND) as send:
        first = client.post(f"{API}/company/invitations/send", headers=auth_headers(admin), json={})
        assert first.json()["enviadas"] == 2 and send.call_count == 2
        again = client.post(f"{API}/company/invitations/send", headers=auth_headers(admin), json={})
        assert again.json()["enviadas"] == 0 and send.call_count == 2  # no se spamea
    assert {m["invite_status"] for e, m in _members(client, auth_headers(admin)).items() if e in (a, b)} == {"sent"}


def test_expired_invitation_can_be_resent_with_new_token(
    client: TestClient, factory, auth_headers
) -> None:
    _org, admin = _setup(factory)
    email = _email()
    _post(client, auth_headers(admin), _book(
        orgs=[["Vencidas", "", 0]],
        members=[["Vencidas", email, "Vence Ya", "collaborator", "", "", ""]],
    ))
    with patch(SEND):
        client.post(f"{API}/company/invitations/send", headers=auth_headers(admin), json={})

    # La invitación vence.
    factory.session.expire_all()
    inv = factory.session.scalar(select(Invitation).where(func.lower(Invitation.email) == email))
    old_hash = inv.token_hash
    inv.expires_at = datetime.now(UTC) - timedelta(days=1)
    factory.session.commit()
    row = _members(client, auth_headers(admin))[email]
    assert row["invite_status"] == "expired"

    with patch(SEND) as send:
        res = client.post(f"{API}/company/members/{row['id']}/resend-invite", headers=auth_headers(admin))
    assert res.status_code == 204, res.text
    send.assert_called_once()
    assert _members(client, auth_headers(admin))[email]["invite_status"] == "sent"

    factory.session.expire_all()
    pending = factory.session.scalars(
        select(Invitation).where(
            func.lower(Invitation.email) == email, Invitation.revoked_at.is_(None),
            Invitation.accepted_at.is_(None),
        )
    ).all()
    assert len(pending) == 1 and pending[0].token_hash != old_hash


def test_resend_rejects_active_members_and_discard_removes_invited(
    client: TestClient, factory, auth_headers
) -> None:
    org, admin = _setup(factory)
    active = factory.make_user(org=org)
    assert (
        client.post(f"{API}/company/members/{active.id}/resend-invite", headers=auth_headers(admin)).status_code
        == 409
    )
    email = _email()
    _post(client, auth_headers(admin), _book(
        orgs=[["Descarte", "", 0]],
        members=[["Descarte", email, "Error De Carga", "collaborator", "", "", ""]],
    ))
    uid = _members(client, auth_headers(admin))[email]["id"]
    assert client.delete(f"{API}/company/members/{uid}/invite", headers=auth_headers(admin)).status_code == 204
    assert email not in _members(client, auth_headers(admin))
    # un activo NO se descarta por esta vía
    assert client.delete(f"{API}/company/members/{active.id}/invite", headers=auth_headers(admin)).status_code == 409


def test_moving_invited_to_another_org_revokes_pending_invitation(
    client: TestClient, factory, auth_headers
) -> None:
    _org, admin = _setup(factory)
    email = _email()
    _post(client, auth_headers(admin), _book(
        orgs=[["Origen", "", 0], ["Destino", "", 0]],
        members=[["Origen", email, "Movido", "collaborator", "", "", ""]],
    ))
    with patch(SEND):
        client.post(f"{API}/company/invitations/send", headers=auth_headers(admin), json={})
    assert _members(client, auth_headers(admin))[email]["invite_status"] == "sent"

    _post(client, auth_headers(admin), _book(members=[["Destino", email, "Movido", "collaborator", "", "", ""]]))
    row = _members(client, auth_headers(admin))[email]
    assert row["org_name"] == "Destino"
    assert row["invite_status"] == "staged"  # el enlace viejo apuntaba a Origen: se revocó


def test_single_sheet_legacy_file_still_works(client: TestClient, factory, auth_headers) -> None:
    org, admin = _setup(factory)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["organizacion", "email", "nombre_completo", "rol", "manager_email"])
    ws.append([org.slug, _email(), "Legado", "collaborator", ""])
    buf = BytesIO()
    wb.save(buf)
    body = _post(client, auth_headers(admin), buf.getvalue()).json()
    assert body["creados"] == 1 and body["errores"] == 0
